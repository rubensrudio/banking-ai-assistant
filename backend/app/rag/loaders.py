from abc import ABC, abstractmethod
from pathlib import Path

from app.core.logger import get_logger

logger = get_logger(__name__)


class BaseDocumentLoader(ABC):
    @abstractmethod
    def load(self, file_path: Path) -> list[dict]:
        """
        Load a document and return a list of page dicts.
        Each dict: {"text": str, "page": int | None, "source": str}
        """


class PDFLoader(BaseDocumentLoader):
    def load(self, file_path: Path) -> list[dict]:
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise ImportError("pypdf is required: pip install pypdf") from e

        reader = PdfReader(str(file_path))

        # Detect repeated header/footer lines across pages (appear on 50%+ of pages)
        all_lines: list[list[str]] = []
        raw_texts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            raw_texts.append(text)
            lines = [l.strip() for l in text.split("\n") if l.strip()]
            all_lines.append(lines)

        repeated_headers = _detect_repeated_lines(all_lines, min_frequency=0.5)

        pages = []
        for i, text in enumerate(raw_texts, start=1):
            pages.append({
                "text": text,
                "page": i,
                "source": file_path.name,
                "_strip_headers": repeated_headers,
            })

        logger.debug(
            f"PDF loaded: {file_path.name} — {len(pages)} pages, "
            f"{len(repeated_headers)} repeated headers detected"
        )
        return pages


def _detect_repeated_lines(all_lines: list[list[str]], min_frequency: float = 0.5) -> list[str]:
    """
    Find lines that appear on at least min_frequency fraction of pages.
    These are likely headers/footers added by the PDF generation tool.
    """
    from collections import Counter

    num_pages = len(all_lines)
    if num_pages < 2:
        return []

    counts: Counter = Counter()
    for lines in all_lines:
        for line in set(lines):  # deduplicate within a page
            counts[line] += 1

    threshold = max(2, int(num_pages * min_frequency))
    return [line for line, count in counts.items() if count >= threshold]


class DOCXLoader(BaseDocumentLoader):
    def load(self, file_path: Path) -> list[dict]:
        try:
            import docx
        except ImportError as e:
            raise ImportError("python-docx is required: pip install python-docx") from e

        doc = docx.Document(str(file_path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        logger.debug(f"DOCX loaded: {file_path.name} — {len(doc.paragraphs)} paragraphs")
        return [{"text": text, "page": None, "source": file_path.name}]


class TXTLoader(BaseDocumentLoader):
    def load(self, file_path: Path) -> list[dict]:
        text = file_path.read_text(encoding="utf-8", errors="replace")
        logger.debug(f"TXT loaded: {file_path.name} — {len(text)} chars")
        return [{"text": text, "page": None, "source": file_path.name}]


class DocumentLoaderFactory:
    _registry: dict[str, type[BaseDocumentLoader]] = {
        ".pdf": PDFLoader,
        ".docx": DOCXLoader,
        ".txt": TXTLoader,
    }

    @classmethod
    def get_loader(cls, file_path: Path) -> BaseDocumentLoader:
        ext = file_path.suffix.lower()
        loader_cls = cls._registry.get(ext)
        if not loader_cls:
            raise ValueError(
                f"Unsupported file type: '{ext}'. "
                f"Supported: {list(cls._registry.keys())}"
            )
        return loader_cls()

    @classmethod
    def supported_extensions(cls) -> list[str]:
        return list(cls._registry.keys())
